"""
Modul inti sistem prediksi cuaca (dipakai oleh dashboard.py dan 08_update_forecast.py).

Alur "terus memperbarui":
  data baru (CSV) -> ditambahkan ke SEAS5_GARAM_2021_2025_AGREGAT.csv
  -> model Transformer membaca 36 dasarian TERAKHIR -> forecast 18 dasarian (6 bulan)
  -> hasil disimpan ke CSV + Excel -> dashboard otomatis menampilkan yang terbaru.
"""
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd

BASE_DIR = Path(__file__).resolve().parents[1]

RAW_FILE = BASE_DIR / "data" / "processed" / "SEAS5_GARAM_2021_2025_AGREGAT.csv"
MODEL_FILE = BASE_DIR / "models" / "transformer_weather_model.keras"
SCALER_FILE = BASE_DIR / "models" / "scaler.pkl"
ENCODER_FILE = BASE_DIR / "models" / "location_encoder.pkl"
EVAL_FILE = BASE_DIR / "results" / "evaluasi" / "evaluasi_transformer_dasarian.json"

FORECAST_CSV = BASE_DIR / "results" / "forecast_cuaca_4lokasi_dasarian_6bulan.csv"
EXCEL_FILE = BASE_DIR / "results" / "prediksi_parameter_cuaca.xlsx"

FEATURES = ["tp_dasarian_mm", "t2m", "d2m", "ssrd_rate", "si10"]
TIME_FEATURES = ["bulan_sin", "bulan_cos", "dasarian_sin", "dasarian_cos"]
KOLOM_WAJIB = ["tahun", "bulan", "dasarian", "lokasi"] + FEATURES
WINDOW = 36      # 12 bulan histori
FORECAST = 18    # 6 bulan prediksi

NAMA_BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni",
              "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
ROMAWI = {1: "I", 2: "II", 3: "III"}


# ---------------------------------------------------------------- data

def _indeks_waktu(df):
    """Nomor urut dasarian (36 per tahun) -- dipakai untuk cek kontinuitas."""
    return df["tahun"] * 36 + (df["bulan"] - 1) * 3 + (df["dasarian"] - 1)


def validasi(df):
    """Pastikan data bersih & tidak ada dasarian yang bolong per lokasi.
    Window 36 dasarian mengasumsikan data berurutan tanpa celah."""
    hilang = [k for k in KOLOM_WAJIB if k not in df.columns]
    if hilang:
        raise ValueError(f"Kolom wajib tidak ditemukan: {hilang}")
    if df[KOLOM_WAJIB].isna().any().any():
        raise ValueError("Ada nilai kosong (NaN) di data.")
    if not df["bulan"].between(1, 12).all() or not df["dasarian"].between(1, 3).all():
        raise ValueError("Nilai bulan harus 1-12 dan dasarian harus 1-3.")

    for lokasi, g in df.groupby("lokasi"):
        selisih = _indeks_waktu(g.sort_values(["tahun", "bulan", "dasarian"])).diff().dropna()
        if (selisih != 1).any():
            raise ValueError(
                f"Data lokasi '{lokasi}' tidak berurutan / ada dasarian yang bolong. "
                "Lengkapi dulu data-nya (tiap bulan harus ada dasarian 1, 2, 3)."
            )
    return df


def load_history():
    df = pd.read_csv(RAW_FILE)
    df = df.sort_values(["lokasi", "tahun", "bulan", "dasarian"]).reset_index(drop=True)
    return df


def tambah_data_baru(df_baru):
    """Gabungkan data baru ke file histori. Baris dengan periode yang sama
    ditimpa dengan nilai terbaru. Mengembalikan jumlah baris baru yang masuk."""
    df_baru = df_baru[KOLOM_WAJIB].copy()
    lama = load_history()[KOLOM_WAJIB]

    gabung = pd.concat([lama, df_baru], ignore_index=True)
    gabung = gabung.drop_duplicates(
        subset=["tahun", "bulan", "dasarian", "lokasi"], keep="last"
    ).sort_values(["lokasi", "tahun", "bulan", "dasarian"]).reset_index(drop=True)

    validasi(gabung)                       # gagal -> file lama tidak disentuh
    gabung.to_csv(RAW_FILE, index=False)
    return len(gabung) - len(lama)


# ---------------------------------------------------------------- celah data & pembaruan otomatis

def _dari_indeks(i):
    t, r = divmod(int(i), 36)
    return t, r // 3 + 1, r % 3 + 1


def isi_celah(df, today=None):
    """Prediksi SELALU mulai bulan depan. Kalau data aktual (ERP) belum sampai
    akhir bulan ini, dasarian yang kosong diisi sementara dengan rata-rata
    historis dasarian yang sama. Begitu data aktual masuk, isian ini hilang
    dengan sendirinya. Return: (data gabungan, hanya baris isian)."""
    now = pd.Timestamp.now() if today is None else pd.Timestamp(today)
    target = now.year * 36 + (now.month - 1) * 3 + 2          # dasarian III bulan ini

    klim = df.groupby(["lokasi", "bulan", "dasarian"])[FEATURES].mean()
    baris = []
    for lokasi, g in df.groupby("lokasi"):
        for j in range(int(_indeks_waktu(g).max()) + 1, target + 1):
            t, b, d = _dari_indeks(j)
            baris.append({"tahun": t, "bulan": b, "dasarian": d, "lokasi": lokasi,
                          **klim.loc[(lokasi, b, d)].to_dict()})
    celah = pd.DataFrame(baris, columns=KOLOM_WAJIB)
    gabung = pd.concat([df[KOLOM_WAJIB], celah], ignore_index=True)
    gabung = gabung.sort_values(["lokasi", "tahun", "bulan", "dasarian"]).reset_index(drop=True)
    return gabung, celah


def awal_forecast():
    """{lokasi: (tahun, bulan, dasarian)} periode pertama prediksi saat ini."""
    gabung, _ = isi_celah(validasi(load_history()))
    hasil = {}
    for lokasi, g in gabung.groupby("lokasi"):
        last = g.iloc[-1]
        hasil[lokasi] = periode_berikutnya(int(last["tahun"]), int(last["bulan"]),
                                           int(last["dasarian"]), 1)[0]
    return hasil


def forecast_usang():
    """True kalau prediksi perlu dihitung ulang: belum ada, data histori lebih
    baru dari prediksi (mis. ERP baru menambah data), atau bulan sudah berganti."""
    if not FORECAST_CSV.exists():
        return True
    if RAW_FILE.stat().st_mtime > FORECAST_CSV.stat().st_mtime:
        return True
    fc = pd.read_csv(FORECAST_CSV)
    for lokasi, awal in awal_forecast().items():
        g = fc[fc["lokasi"] == lokasi]
        if g.empty or tuple(int(v) for v in g.iloc[0][["tahun", "bulan", "dasarian"]]) != awal:
            return True
    return False


# ---------------------------------------------------------------- forecast

def periode_berikutnya(tahun, bulan, dasarian, n):
    hasil, t, b, d = [], tahun, bulan, dasarian
    for _ in range(n):
        d += 1
        if d > 3:
            d, b = 1, b + 1
            if b > 12:
                b, t = 1, t + 1
        hasil.append((t, b, d))
    return hasil


def jalankan_forecast(simpan=True):
    """Forecast 18 dasarian ke depan dari data terakhir untuk semua lokasi."""
    import joblib
    from tensorflow.keras.models import load_model   # import berat -> lazy

    model = load_model(MODEL_FILE)
    scaler = joblib.load(SCALER_FILE)
    encoder = joblib.load(ENCODER_FILE)

    df, _ = isi_celah(validasi(load_history()))
    hasil = []

    for lokasi in encoder.classes_:
        g = df[df["lokasi"] == lokasi].reset_index(drop=True)
        if len(g) < WINDOW:
            raise ValueError(f"Data '{lokasi}' kurang dari {WINDOW} dasarian.")

        # skala memakai scaler LAMA (sama seperti saat training) -> tidak di-fit ulang
        g = g.copy()
        g[FEATURES] = scaler.transform(g[FEATURES])
        g["bulan_sin"] = np.sin(2 * np.pi * g["bulan"] / 12)
        g["bulan_cos"] = np.cos(2 * np.pi * g["bulan"] / 12)
        g["dasarian_sin"] = np.sin(2 * np.pi * g["dasarian"] / 3)
        g["dasarian_cos"] = np.cos(2 * np.pi * g["dasarian"] / 3)

        x_seq = g[FEATURES + TIME_FEATURES].tail(WINDOW).to_numpy("float32")[None]
        x_loc = np.array([[encoder.transform([lokasi])[0]]], dtype="int32")

        pred = model.predict([x_seq, x_loc], verbose=0)[0]          # (18, 5)
        pred = scaler.inverse_transform(pred)
        pred[:, 0] = np.clip(pred[:, 0], 0, None)                   # hujan tidak negatif

        last = g.iloc[-1]
        periode = periode_berikutnya(int(last["tahun"]), int(last["bulan"]),
                                     int(last["dasarian"]), FORECAST)
        f = pd.DataFrame(pred, columns=FEATURES)
        f.insert(0, "lokasi", lokasi)
        f.insert(1, "tahun", [p[0] for p in periode])
        f.insert(2, "bulan", [p[1] for p in periode])
        f.insert(3, "dasarian", [p[2] for p in periode])
        f.insert(4, "Periode", [f"{NAMA_BULAN[b-1]} {t} Dasarian {ROMAWI[d]}"
                                for t, b, d in periode])
        hasil.append(f)

    forecast = pd.concat(hasil, ignore_index=True)

    if simpan:
        FORECAST_CSV.parent.mkdir(exist_ok=True)
        forecast.to_csv(FORECAST_CSV, index=False)
        EXCEL_FILE.write_bytes(buat_excel(forecast))
    return forecast


# ---------------------------------------------------------------- excel

KOLOM_EXCEL = {
    "Periode": "Periode",
    "tp_dasarian_mm": "Curah Hujan (mm)",
    "Suhu": "Suhu Udara (°C)",
    "Embun": "Titik Embun (°C)",
    "ssrd_rate": "Radiasi Matahari (W/m²)",
    "si10": "Kecepatan Angin (m/s)",
}


def _tabel_excel(f):
    t = pd.DataFrame({
        "Periode": f["Periode"],
        "Tahun": f["tahun"], "Bulan": f["bulan"], "Dasarian": f["dasarian"],
        KOLOM_EXCEL["Suhu"]: (f["t2m"] - 273.15).round(2),
        KOLOM_EXCEL["Embun"]: (f["d2m"] - 273.15).round(2),
        KOLOM_EXCEL["ssrd_rate"]: f["ssrd_rate"].round(2),
        KOLOM_EXCEL["si10"]: f["si10"].round(2),
        KOLOM_EXCEL["tp_dasarian_mm"]: f["tp_dasarian_mm"].round(1),
    })
    return t


def buat_excel(forecast):
    """Excel: 1 sheet 'Semua Lokasi' + 1 sheet per lokasi. Return bytes."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        semua = _tabel_excel(forecast)
        semua.insert(0, "Lokasi", forecast["lokasi"].values)
        semua.to_excel(xw, sheet_name="Semua Lokasi", index=False)
        for lokasi, g in forecast.groupby("lokasi", sort=True):
            _tabel_excel(g).to_excel(xw, sheet_name=lokasi[:31], index=False)

        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="2563EB")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[1].height = 32
            ws.freeze_panes = "A2"
            for i, col in enumerate(ws.columns, 1):
                lebar = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(lebar + 3, 11), 28)
    return buf.getvalue()
